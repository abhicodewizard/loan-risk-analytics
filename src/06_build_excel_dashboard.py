"""
Step 6 - Build the Excel credit-risk dashboard.

* 'Data' sheet : 30,000 applications (risk bands and loss estimates come from the SQL view vw_loans)
* every KPI / table on the other sheets is a LIVE Excel formula over 'Data'
  (COUNTIFS, SUMIFS, AVERAGEIFS, RANK) - change the data and the dashboard updates.

Run:  python src/06_build_excel_dashboard.py     (then open in Excel; formulas calculate on load)
"""
import sqlite3
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "excel" / "Loan_Risk_Dashboard.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

conn = sqlite3.connect(ROOT / "database" / "loans.db")
df = pd.read_sql("""SELECT application_id, CAST(strftime('%Y%m', application_date) AS INTEGER) AS app_yyyymm, city, region, product,
       employment_type, channel, score_band,
       CASE WHEN foir < .30 THEN '1. Below 30%' WHEN foir < .40 THEN '2. 30-40%' WHEN foir < .50 THEN '3. 40-50%'
            WHEN foir < .60 THEN '4. 50-60%' ELSE '5. 60%+' END AS foir_band,
       decision, is_disbursed, is_npa, is_delinquent, loan_amount, outstanding_principal, interest_rate, mob, loan_status,
       credit_loss_est, interest_earned, processing_fee
       FROM vw_loans ORDER BY application_id""", conn)
N = len(df) + 1
months = sorted(df.app_yyyymm.unique())
cities = pd.read_sql("""SELECT city FROM vw_loans WHERE is_disbursed=1 AND mob>=12 GROUP BY city ORDER BY 1.0*SUM(is_npa)/COUNT(*) DESC""", conn).city.tolist()
products = pd.read_sql("""SELECT product FROM vw_loans WHERE is_disbursed=1 GROUP BY product ORDER BY SUM(loan_amount) DESC""", conn)["product"].tolist()
bands = sorted(df.score_band.unique())
foirs = sorted(df.foir_band.unique())

FONT, NAVY, LIGHT = "Arial", "1F4E78", "DDEBF7"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FILL = PatternFill("solid", fgColor=NAVY)


def hdr(ws, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.font, c.fill, c.border = Font(name=FONT, bold=True, color="FFFFFF"), H_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def fmt_range(ws, r1, r2, formats):
    for r in range(r1, r2 + 1):
        for c, f in formats.items():
            ws.cell(r, c).number_format = f
        for c in range(1, max(formats) + 1):
            ws.cell(r, c).font, ws.cell(r, c).border = Font(name=FONT), BORDER


wb = Workbook()
dash, mon, seg, prd = wb.active, wb.create_sheet("Monthly"), wb.create_sheet("Segments"), wb.create_sheet("Products")
dash.title = "Dashboard"
data = wb.create_sheet("Data")

cols = list(df.columns)
for i, c in enumerate(cols, 1):
    cell = data.cell(1, i, c.replace("_", " ").title())
    cell.font, cell.fill = Font(name=FONT, bold=True, color="FFFFFF"), H_FILL
    data.column_dimensions[get_column_letter(i)].width = 15
for row in df.astype(object).where(df.notna(), None).itertuples(index=False):
    data.append(list(row))
data.freeze_panes = "A2"
data.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{N}"
L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
R = lambda c: f"Data!${L[c]}$2:${L[c]}${N}"
MATURE = f'{R("is_disbursed")},1,{R("mob")},">=12"'            # loans with 12+ months on book

# ---------------- KPI cards ----------------
dash.sheet_view.showGridLines = False
dash["A1"] = "Loan Portfolio & Credit Risk Dashboard"
dash["A1"].font = Font(name=FONT, size=20, bold=True, color=NAVY)
dash["A2"] = ("Simulated retail-lending portfolio (30,000 applications, Jan 2023 - Sep 2025, as of 31-Dec-2025) | every number is a live formula over the "
              "Data sheet | risk bands and loss estimates come from the SQL view vw_loans")
dash["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
OUTS = f'SUMIFS({R("outstanding_principal")},{R("is_disbursed")},1,{R("loan_status")},"<>Closed")'
kpis = [
    ("Applications", f"=COUNTA({R('application_id')})", "#,##0"),
    ("Approval rate", f"=SUM({R('is_disbursed')})/A5", "0.0%"),
    ("Disbursed (Rs Cr)", f'=SUMIFS({R("loan_amount")},{R("is_disbursed")},1)/10^7', "#,##0.0"),
    ("Outstanding book (Rs Cr)", f"={OUTS}/10^7", "#,##0.0"),
    ("Gross NPA %", f'=SUMIFS({R("outstanding_principal")},{R("is_npa")},1)/({OUTS})', "0.00%"),
    ("30-89 DPD (SMA) %", f'=SUMIFS({R("outstanding_principal")},{R("is_delinquent")},1)/({OUTS})', "0.00%"),
    ("Est. credit loss (Rs Cr)", f"=SUM({R('credit_loss_est')})/10^7", "#,##0.0"),
    ("Avg interest rate", f'=AVERAGEIFS({R("interest_rate")},{R("is_disbursed")},1)/100', "0.00%"),
]
for i, (lab, f, fm) in enumerate(kpis, 1):
    a = dash.cell(4, i, lab)
    a.font, a.fill = Font(name=FONT, bold=True, color="FFFFFF", size=10), H_FILL
    a.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    b = dash.cell(5, i, f)
    b.font, b.fill = Font(name=FONT, bold=True, size=15, color=NAVY), PatternFill("solid", fgColor=LIGHT)
    b.alignment, b.number_format = Alignment(horizontal="center", vertical="center"), fm
    a.border = b.border = BORDER
    dash.column_dimensions[get_column_letter(i)].width = 19
dash.row_dimensions[4].height, dash.row_dimensions[5].height = 30, 34

# ---------------- Monthly ----------------
mon["A1"] = "Monthly application funnel"
mon["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
hdr(mon, 3, ["Month (YYYYMM)", "Applications", "Approved", "Approval %", "Disbursed (Rs Cr)"])
for i, ym in enumerate(months):
    r = 4 + i
    mon.cell(r, 1, int(ym))
    mon.cell(r, 2, f"=COUNTIFS({R('app_yyyymm')},$A{r})")
    mon.cell(r, 3, f"=SUMIFS({R('is_disbursed')},{R('app_yyyymm')},$A{r})")
    mon.cell(r, 4, f"=IFERROR(C{r}/B{r},0)")
    mon.cell(r, 5, f"=SUMIFS({R('loan_amount')},{R('app_yyyymm')},$A{r},{R('is_disbursed')},1)/10^7")
last_m = 3 + len(months)
fmt_range(mon, 4, last_m, {1: "0", 2: "#,##0", 3: "#,##0", 4: "0.0%", 5: "#,##0.00"})
for c in "ABCDE":
    mon.column_dimensions[c].width = 17

# ---------------- Segments ----------------
seg["A1"] = "Risk by segment  (NPA rate measured on loans with 12+ months on book)"
seg["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
BLOCKS = {}


def block(top, title, key, values):
    seg.cell(top, 1, title).font = Font(name=FONT, bold=True, color=NAVY)
    hdr(seg, top + 1, [title.split(" by ")[-1].title(), "Applications", "Approved", "Approval %", "Matured loans", "NPA loans", "NPA rate"])
    for i, v in enumerate(values):
        r = top + 2 + i
        seg.cell(r, 1, v)
        seg.cell(r, 2, f"=COUNTIFS({R(key)},$A{r})")
        seg.cell(r, 3, f"=SUMIFS({R('is_disbursed')},{R(key)},$A{r})")
        seg.cell(r, 4, f"=IFERROR(C{r}/B{r},0)")
        seg.cell(r, 5, f"=COUNTIFS({R(key)},$A{r},{MATURE})")
        seg.cell(r, 6, f"=COUNTIFS({R(key)},$A{r},{MATURE},{R('is_npa')},1)")
        seg.cell(r, 7, f"=IFERROR(F{r}/E{r},0)")
    r1, r2 = top + 2, top + 1 + len(values)
    fmt_range(seg, r1, r2, {1: "@", 2: "#,##0", 3: "#,##0", 4: "0.0%", 5: "#,##0", 6: "#,##0", 7: "0.0%"})
    seg.conditional_formatting.add(f"G{r1}:G{r2}", ColorScaleRule(start_type="min", start_color="C6E0B4", mid_type="percentile", mid_value=50,
                                   mid_color="FFEB84", end_type="max", end_color="F8696B"))
    BLOCKS[key] = (top, r1, r2)
    return r2 + 3


nxt = block(3, "Risk by credit score band", "score_band", bands)
nxt = block(nxt, "Risk by FOIR band", "foir_band", foirs)
nxt = block(nxt, "Risk by channel", "channel", ["Branch", "Online", "DSA"])
nxt = block(nxt, "Risk by employment type", "employment_type", ["Salaried", "Self-Employed", "Business Owner"])
nxt = block(nxt, "Risk by city", "city", cities)
c1, c2 = BLOCKS["city"][1], BLOCKS["city"][2]
seg.cell(BLOCKS["city"][0] + 1, 8, "Risk rank").font = Font(name=FONT, bold=True, color="FFFFFF")
seg.cell(BLOCKS["city"][0] + 1, 8).fill, seg.cell(BLOCKS["city"][0] + 1, 8).border = H_FILL, BORDER
for r in range(c1, c2 + 1):
    seg.cell(r, 8, f"=RANK(G{r},$G${c1}:$G${c2})").border = BORDER
    seg.cell(r, 8).font = Font(name=FONT)
for c, w in zip("ABCDEFGH", [20, 13, 12, 12, 14, 11, 11, 10]):
    seg.column_dimensions[c].width = w

# ---------------- Products ----------------
prd["A1"] = "Product risk and risk-adjusted return"
prd["A1"].font = Font(name=FONT, bold=True, size=13, color=NAVY)
hdr(prd, 3, ["Product", "Disbursed loans", "Disbursed (Rs Cr)", "Matured loans", "NPA loans", "NPA rate", "Interest earned (Rs Cr)",
             "Processing fees (Rs Cr)", "Credit loss (Rs Cr)", "Risk-adj. income (Rs Cr)", "Return on disbursed %"])
for i, p in enumerate(products):
    r = 4 + i
    prd.cell(r, 1, p)
    prd.cell(r, 2, f"=COUNTIFS({R('product')},$A{r},{R('is_disbursed')},1)")
    prd.cell(r, 3, f"=SUMIFS({R('loan_amount')},{R('product')},$A{r},{R('is_disbursed')},1)/10^7")
    prd.cell(r, 4, f"=COUNTIFS({R('product')},$A{r},{MATURE})")
    prd.cell(r, 5, f"=COUNTIFS({R('product')},$A{r},{MATURE},{R('is_npa')},1)")
    prd.cell(r, 6, f"=IFERROR(E{r}/D{r},0)")
    prd.cell(r, 7, f"=SUMIFS({R('interest_earned')},{R('product')},$A{r})/10^7")
    prd.cell(r, 8, f"=SUMIFS({R('processing_fee')},{R('product')},$A{r})/10^7")
    prd.cell(r, 9, f"=SUMIFS({R('credit_loss_est')},{R('product')},$A{r})/10^7")
    prd.cell(r, 10, f"=G{r}+H{r}-I{r}")
    prd.cell(r, 11, f"=IFERROR(J{r}/C{r},0)")
last_p = 3 + len(products)
fmt_range(prd, 4, last_p, {1: "@", 2: "#,##0", 3: "#,##0.0", 4: "#,##0", 5: "#,##0", 6: "0.0%", 7: "#,##0.00", 8: "#,##0.00", 9: "#,##0.00", 10: "#,##0.00", 11: "0.0%"})
for c in range(1, 12):
    prd.column_dimensions[get_column_letter(c)].width = 16
prd.column_dimensions["A"].width = 18

# ---------------- charts ----------------
def style_axes(ch):
    ch.x_axis.delete = ch.y_axis.delete = False


bar = BarChart()
bar.type, bar.title, bar.style = "col", "Monthly applications vs approval %", 10
bar.add_data(Reference(mon, min_col=2, min_row=3, max_row=last_m), titles_from_data=True)
bar.set_categories(Reference(mon, min_col=1, min_row=4, max_row=last_m))
ln = LineChart()
ln.add_data(Reference(mon, min_col=4, min_row=3, max_row=last_m), titles_from_data=True)
ln.series[0].smooth = False
ln.y_axis.axId, ln.y_axis.crosses, ln.y_axis.number_format = 200, "max", "0%"
ln.y_axis.scaling.min, ln.y_axis.scaling.max = 0.4, 0.8
bar += ln
bar.y_axis.scaling.min = 0
bar.height, bar.width = 8.5, 17
style_axes(bar)
ln.y_axis.delete = False
dash.add_chart(bar, "A7")

s0, s1, s2 = BLOCKS["score_band"]
sc = BarChart()
sc.type, sc.title, sc.style = "col", "NPA rate by credit score band", 10
sc.add_data(Reference(seg, min_col=7, min_row=s0 + 1, max_row=s2), titles_from_data=True)
sc.set_categories(Reference(seg, min_col=1, min_row=s1, max_row=s2))
sc.legend, sc.height, sc.width = None, 8.5, 15
sc.y_axis.number_format = "0%"
sc.dataLabels = DataLabelList()
sc.dataLabels.showVal, sc.dataLabels.showSerName, sc.dataLabels.showCatName, sc.dataLabels.showLegendKey = True, False, False, False
style_axes(sc)
dash.add_chart(sc, "F7")

pn = BarChart()
pn.type, pn.title, pn.style = "bar", "NPA rate by product", 10
pn.add_data(Reference(prd, min_col=6, min_row=3, max_row=last_p), titles_from_data=True)
pn.set_categories(Reference(prd, min_col=1, min_row=4, max_row=last_p))
pn.legend, pn.height, pn.width = None, 8, 12
pn.y_axis.number_format = "0%"
pn.y_axis.scaling.min = 0
pn.dataLabels = DataLabelList()
pn.dataLabels.showVal, pn.dataLabels.showSerName, pn.dataLabels.showCatName, pn.dataLabels.showLegendKey = True, False, False, False
style_axes(pn)
dash.add_chart(pn, "A25")

rt = BarChart()
rt.type, rt.title, rt.style = "bar", "Risk-adjusted return on disbursed % by product", 10
rt.add_data(Reference(prd, min_col=11, min_row=3, max_row=last_p), titles_from_data=True)
rt.set_categories(Reference(prd, min_col=1, min_row=4, max_row=last_p))
rt.legend, rt.height, rt.width = None, 8, 12
rt.y_axis.number_format = "0%"
rt.y_axis.scaling.min = 0
rt.dataLabels = DataLabelList()
rt.dataLabels.showVal, rt.dataLabels.showSerName, rt.dataLabels.showCatName, rt.dataLabels.showLegendKey = True, False, False, False
style_axes(rt)
dash.add_chart(rt, "F25")

dash["A42"] = ("Definitions: NPA = 90+ days past due. Gross NPA % = NPA outstanding / live book outstanding. NPA rate in segment tables = NPA loans / loans with 12+ months on "
               "book (so segments are comparable). Credit loss = NPA outstanding x product LGD. Risk-adjusted income = interest earned + 1% processing fee - credit loss.")
dash["A42"].font = Font(name=FONT, italic=True, size=9, color="595959")
dash["A42"].alignment = Alignment(wrap_text=True, vertical="top")
dash.merge_cells("A42:I43")
dash.page_setup.orientation = "landscape"
dash.page_setup.fitToWidth, dash.page_setup.fitToHeight = 1, 0
dash.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("Saved", OUT.relative_to(ROOT))
